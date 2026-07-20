import docx
import os
import json

def fix_template_1(path):
    if not os.path.exists(path):
        return "Not found"
        
    doc = docx.Document(path)
    modified = False
    
    for p in doc.paragraphs:
        # Fix the parenthesized total price block
        if "ZAMON_BUYURTMACHI_FISH" in p.text and "сум" in p.text and "(" in p.text:
            # We want to replace this paragraph's text with: ( {{ ZAMON_UMUMIY_NARX_SOZ }} ) сум.
            p.text = "( {{ ZAMON_UMUMIY_NARX_SOZ }} ) сум."
            modified = True
            
        # Clean trailing underscores after buyer F.I.Sh
        if "{{ ZAMON_BUYURTMACHI_FISH }}___________" in p.text:
            p.text = p.text.replace("{{ ZAMON_BUYURTMACHI_FISH }}___________", "{{ ZAMON_BUYURTMACHI_FISH }}")
            modified = True
            
        # Set page break before Shartnoma header
        if "ША Р Т Н О М А" in p.text:
            if not p.paragraph_format.page_break_before:
                p.paragraph_format.page_break_before = True
                modified = True
                
    if modified:
        doc.save(path)
        return "Template 1 final fixes applied"
    return "No changes in Template 1"


def fix_template_2(path):
    if not os.path.exists(path):
        return "Not found"
        
    doc = docx.Document(path)
    modified = False
    
    for p in doc.paragraphs:
        # Law Date
        if "{{ MAYDON }}.1998" in p.text:
            p.text = p.text.replace("{{ MAYDON }}.1998", "29.08.1998")
            modified = True
            
        # Set page break before Shartnoma header
        if "ША Р Т Н О М А" in p.text:
            if not p.paragraph_format.page_break_before:
                p.paragraph_format.page_break_before = True
                modified = True
                
    if modified:
        doc.save(path)
        return "Template 2 final fixes applied"
    return "No changes in Template 2"

p1 = '/var/www/sellerarp/storage/app/private/contract_templates/tpl_1782948193_6a45a161f4175.docx'
p2 = '/var/www/sellerarp/storage/app/private/contract_templates/tpl_1783451232_6a4d4e60c2bc7.docx'

print(json.dumps([fix_template_1(p1), fix_template_2(p2)]))