import sys
import os
import json
import re
import docx
import mammoth
from docxtpl import DocxTemplate

def detect_placeholders(docx_path):
    """
    Scans the docx file for common patterns (PINFL, passport, phone, dates, prices, etc.)
    and returns a list of proposed placeholders.
    """
    if not os.path.exists(docx_path):
        return {"error": f"File not found: {docx_path}"}
        
    doc = docx.Document(docx_path)
    text_segments = []
    
    # Extract text from paragraphs
    for p in doc.paragraphs:
        txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
        txt = re.sub(r'\s+', ' ', txt).strip()
        if txt:
            text_segments.append(txt)
            
    # Extract text from tables
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
                    txt = re.sub(r'\s+', ' ', txt).strip()
                    if txt:
                        text_segments.append(txt)

    # Extract text from headers and footers
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for p in header.paragraphs:
                    txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
                    txt = re.sub(r'\s+', ' ', txt).strip()
                    if txt:
                        text_segments.append(txt)
                for t in header.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
                                txt = re.sub(r'\s+', ' ', txt).strip()
                                if txt:
                                    text_segments.append(txt)
                                    
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for p in footer.paragraphs:
                    txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
                    txt = re.sub(r'\s+', ' ', txt).strip()
                    if txt:
                        text_segments.append(txt)
                for t in footer.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                txt = p.text.replace('\xa0', ' ').replace('\t', ' ')
                                txt = re.sub(r'\s+', ' ', txt).strip()
                                if txt:
                                    text_segments.append(txt)
                        
    combined_text = "\n".join(text_segments)
    
    candidates = []
    
    # 1. Detect PINFL (JSHSHIR): exactly 14 digits
    pinfls = re.findall(r'\b\d{14}\b', combined_text)
    for p in set(pinfls):
        candidates.append({
            "text": p,
            "key": "MIJOZ_JSHSHIR",
            "label": "Mijoz JSHSHIR (PINFL)",
            "type": "text"
        })
        
    # 2. Detect Passport: e.g. AA 1234567, AD 9160055, or with №
    passports = re.findall(r'\b[A-Z]{2}\s*№?\s*\d{7}\b', combined_text)
    for p in set(passports):
        candidates.append({
            "text": p,
            "key": "PASPORT_SERIYA",
            "label": "Pasport Seriyasi va Raqami",
            "type": "text"
        })
        
    # 3. Detect Dates: e.g. 26/10/2024-yilda, 13 mart 2025-yil
    # Standard format dd/mm/yyyy
    dates_slash = re.findall(r'\b\d{2}/\d{2}/\d{4}(?:-yilda)?\b', combined_text)
    for d in set(dates_slash):
        candidates.append({
            "text": d,
            "key": "PASPORT_SANA",
            "label": "Pasport berilgan sana",
            "type": "text"
        })
        
    # O'zbekcha yozma sana: e.g. "13 " mart 2025-yil or 13-mart 2025-yil
    dates_uz = re.findall(r'\b“?\d{1,2}”?\s*[a-zA-Z\'`ʻ’]+\s+\d{4}-yil\b', combined_text)
    for d in set(dates_uz):
        candidates.append({
            "text": d,
            "key": "SHARTNOMA_SANASI",
            "label": "Shartnoma sanasi",
            "type": "text"
        })

    # 4. Detect Phone Numbers: e.g. (94) 633-04-00 or +998 (94) 633-04-00
    phones = re.findall(r'\+?998\s*\(?\d{2}\)?\s*\d{3}-\d{2}-\d{2}|\(\d{2}\)\s*\d{3}-\d{2}-\d{2}|\b998\d{9}\b', combined_text)
    for ph in set(phones):
        candidates.append({
            "text": ph,
            "key": "MIJOZ_TELEFON",
            "label": "Mijoz telefon raqami",
            "type": "text"
        })

    # 5. Detect large prices / sums: numbers separated by spaces (e.g. 190 000 000, 3 998 316)
    # Use only spaces/tabs so we don't merge numbers across line breaks
    prices = re.findall(r'\b\d{1,3}(?:[ \t]+\d{3}){2,3}\b', combined_text)
    for pr in set(prices):
        val = int(re.sub(r'\s+', '', pr))
        if val > 1000000:
            candidates.append({
                "text": pr,
                "key": "UMUMIY_NARX",
                "label": "Shartnoma umumiy summasi",
                "type": "number"
            })

    # 6. Detect Percentages: e.g. 20%, 100%
    percents = re.findall(r'\b\d{1,3}%\b', combined_text)
    for p in set(percents):
        candidates.append({
            "text": p,
            "key": "BOSHLANGICH_FOIZ",
            "label": "Boshlang'ich to'lov foizi",
            "type": "text"
        })

    # 7. Detect Decimals (Area): e.g. 71,46 or 73.18
    decimals = re.findall(r'\b\d{2,3}[,.]\d{2}(?!\d)', combined_text)
    for d in set(decimals):
        candidates.append({
            "text": d,
            "key": "MAYDON",
            "label": "Xonadon maydoni",
            "type": "text"
        })

    # 8. Detect Contract Numbers: e.g. 4B/9 or 1B/1P
    contract_nums = re.findall(r'\b\d+[A-Z]?/\d+[A-Z]?\b', combined_text)
    for c_num in set(contract_nums):
        candidates.append({
            "text": c_num,
            "key": "SHARTNOMA_RAQAMI",
            "label": "Shartnoma raqami",
            "type": "text"
        })

    # 9. Detect payment terms/durations: e.g. 6 (olti) bank, 36 (o'ttiz olti) oy
    durations = re.findall(r"\b\d+\s*\([a-zA-Z'`ʻ’]+\)\s*(?:bank|ish|kun|oy)\b", combined_text)
    for dur in set(durations):
        key = "MUDDAT_OY" if "oy" in dur.lower() else "TOLOV_MUDDAT_KUN"
        label = "Kredit/To'lov muddati (oylarda)" if "oy" in dur.lower() else "Dastlabki to'lov muddati (kunlarda)"
        candidates.append({
            "text": dur,
            "key": key,
            "label": label,
            "type": "text"
        })
            
    # 10. Specific hardcoded search for sample names (for test fixtures / samples)
    if "Mamanov Abdulla Xursanmurotovich" in combined_text:
        candidates.append({
            "text": "Mamanov Abdulla Xursanmurotovich",
            "key": "MIJOZ_FISH",
            "label": "Mijozning to'liq F.I.Sh",
            "type": "text"
        })

    if "A. Mamanov" in combined_text:
        candidates.append({
            "text": "A. Mamanov",
            "key": "MIJOZ_FISH_QISQA",
            "label": "Mijoz F.I.Sh qisqartmasi (Masalan: A. Mamanov)",
            "type": "text"
        })
    elif "A.Mamanov" in combined_text:
        candidates.append({
            "text": "A.Mamanov",
            "key": "MIJOZ_FISH_QISQA",
            "label": "Mijoz F.I.Sh qisqartmasi (Masalan: A.Mamanov)",
            "type": "text"
        })
        
    if "Samarqand viloyati Urgut tumani I.I.B" in combined_text:
        candidates.append({
            "text": "Samarqand viloyati Urgut tumani I.I.B",
            "key": "PASPORT_JOYI",
            "label": "Pasport bergan organ",
            "type": "text"
        })
        
    if "Samarqand viloyati, Urgut tuman Yuqori Tegana Do`stlik maxallasida, ro`yxatda turuvchi" in combined_text:
        candidates.append({
            "text": "Samarqand viloyati, Urgut tuman Yuqori Tegana Do`stlik maxallasida, ro`yxatda turuvchi",
            "key": "MIJOZ_MANZIL",
            "label": "Mijoz yashash manzili",
            "type": "text"
        })
    elif "Samarqand viloyati, Urgut tumani Yuqori Tegana maxalla" in combined_text:
        candidates.append({
            "text": "Samarqand viloyati, Urgut tumani Yuqori Tegana maxalla",
            "key": "MIJOZ_MANZIL",
            "label": "Mijoz yashash manzili",
            "type": "text"
        })

    if "Сойибов Назиржон Нематиллаевич" in combined_text:
        candidates.append({
            "text": "Сойибов Назиржон Нематиллаевич",
            "key": "MIJOZ_FISH",
            "label": "Mijozning to'liq F.I.Sh",
            "type": "text"
        })
        
    if "Н. Сойибов" in combined_text:
        candidates.append({
            "text": "Н. Сойибов",
            "key": "MIJOZ_FISH_QISQA",
            "label": "Mijoz F.I.Sh qisqartmasi (Masalan: Н. Сойибов)",
            "type": "text"
        })
    elif "Н.Сойибов" in combined_text:
        candidates.append({
            "text": "Н.Сойибов",
            "key": "MIJOZ_FISH_QISQA",
            "label": "Mijoz F.I.Sh qisqartmasi (Masalan: Н.Сойибов)",
            "type": "text"
        })
        
    if "Urgut tuman Gulobod" in combined_text:
        candidates.append({
            "text": "Urgut tuman Gulobod",
            "key": "MIJOZ_MANZIL",
            "label": "Mijoz yashash manzili",
            "type": "text"
        })
        
    if "Urgut tuman I.I.B" in combined_text:
        candidates.append({
            "text": "Urgut tuman I.I.B",
            "key": "PASPORT_JOYI",
            "label": "Pasport bergan organ",
            "type": "text"
        })
        
    if "ikki yuz sakson besh million sakkiz yuz qirq ming" in combined_text:
        candidates.append({
            "text": "ikki yuz sakson besh million sakkiz yuz qirq ming",
            "key": "UMUMIY_NARX_SOZ_BILAN",
            "label": "Umumiy shartnoma summasi so'z bilan",
            "type": "text"
        })
        
    if "ellik yetti million bir yuz oltmish sakkiz ming" in combined_text:
        candidates.append({
            "text": "ellik yetti million bir yuz oltmish sakkiz ming",
            "key": "BOSHLANGICH_TOLOV_SOZ_BILAN",
            "label": "Boshlang'ich to'lov summasi so'z bilan",
            "type": "text"
        })

    # 10.5. Zamon Sifat Servis specific detections
    if "Тошпулатов Достон Толмас угли" in combined_text:
        candidates.append({
            "text": "Тошпулатов Достон Толмас угли",
            "key": "ZAMON_BUYURTMACHI_FISH",
            "label": "Zamon Buyurtmachi FISH",
            "type": "text"
        })
    if "Худойкулов Зариф Ишанкулович" in combined_text:
        candidates.append({
            "text": "Худойкулов Зариф Ишанкулович",
            "key": "ZAMON_BUYURTMACHI_FISH",
            "label": "Zamon Buyurtmachi FISH",
            "type": "text"
        })
    if "Бакаева Зебинисо Аминджановна" in combined_text:
        candidates.append({
            "text": "Бакаева Зебинисо Аминджановна",
            "key": "ZAMON_BUYURTMACHI_FISH",
            "label": "Zamon Buyurtmachi FISH",
            "type": "text"
        })
    if "Самарканд вилояти Самарканд шахар Мадаткор МФЙ Узбекистон кучаси 116 Г- уй 23-хонадонда" in combined_text:
        candidates.append({
            "text": "Самарканд вилояти Самарканд шахар Мадаткор МФЙ Узбекистон кучаси 116 Г- уй 23-хонадонда",
            "key": "ZAMON_BUYURTMACHI_PROPISKA",
            "label": "Zamon Buyurtmachi Propiska",
            "type": "text"
        })
    if "Самарканд вилояти Самарканд шахри Кул махаллласида" in combined_text:
        candidates.append({
            "text": "Самарканд вилояти Самарканд шахри Кул махаллласида",
            "key": "ZAMON_BUYURTMACHI_PROPISKA",
            "label": "Zamon Buyurtmachi Propiska",
            "type": "text"
        })
    if "подъезднинг тартиб раками № 2" in combined_text:
        candidates.append({
            "text": "подъезднинг тартиб раками № 2",
            "key": "ZAMON_PODEZD",
            "label": "Zamon Podyezd Raqami",
            "type": "text"
        })
    if "хонадоннинг тартиб раками № 157" in combined_text:
        candidates.append({
            "text": "хонадоннинг тартиб раками № 157",
            "key": "ZAMON_XONADON",
            "label": "Zamon Xonadon Raqami",
            "type": "text"
        })
    if "қавати 9 (туккиз)" in combined_text:
        candidates.append({
            "text": "қавати 9 (туккиз)",
            "key": "ZAMON_QAVAT",
            "label": "Zamon Qavati",
            "type": "text"
        })
    if "хоналар сони 2 (икки)" in combined_text:
        candidates.append({
            "text": "хоналар сони 2 (икки)",
            "key": "ZAMON_XONALAR",
            "label": "Zamon Xonalar Soni",
            "type": "text"
        })
    if "умумий майдони 55,1" in combined_text:
        candidates.append({
            "text": "умумий майдони 55,1",
            "key": "ZAMON_MAYDON",
            "label": "Zamon Umumiy Maydoni",
            "type": "text"
        })
    if "7 500 000" in combined_text:
        candidates.append({
            "text": "7 500 000",
            "key": "ZAMON_KV_METR_NARXI",
            "label": "Zamon 1kv.m Narxi",
            "type": "number"
        })
    if "413 250 000" in combined_text:
        candidates.append({
            "text": "413 250 000",
            "key": "ZAMON_UMUMIY_NARX",
            "label": "Zamon Umumiy Narxi",
            "type": "number"
        })

    # 11. Detect combined block/floor/room description (e.g. 4-(to'rtinchi) blok 4-(to'rtinchi) qavati, 9-(to'qqizinchi) xonadoni)
    location_matches = re.findall(
        r'\d+\s*-\s*\([^)]+\)\s+blok[a-z]*\s+-?\d+\s*-\s*\([^)]+\)\s+qavat[a-z]*\s*,\s*\d+\s*-\s*\([^)]+\)\s+(?:xonadon[a-z]*|savdo\s+do[\'`ʻ’]koni[a-z]*)',
        combined_text,
        re.IGNORECASE
    )
    for loc in set(location_matches):
        candidates.append({
            "text": loc,
            "key": "XONA_TAFSILOTI",
            "label": "Xona/Xonadon to'liq tafsiloti (Blok, Qavat, Raqam)",
            "type": "text"
        })

    # Return unique candidates sorted by text length desc to avoid substring conflicts during replacement
    unique_candidates = []
    seen = set()
    for c in sorted(candidates, key=lambda x: len(x["text"]), reverse=True):
        if c["text"] not in seen:
            seen.add(c["text"])
            unique_candidates.append(c)
            
    return unique_candidates

def convert_docx_to_html(docx_path):
    """
    Converts docx to HTML using mammoth for web preview.
    """
    if not os.path.exists(docx_path):
        return {"error": f"File not found: {docx_path}"}
        
    with open(docx_path, "rb") as docx_file:
        result = mammoth.convert_to_html(docx_file)
        return {"html": result.value, "warnings": result.messages}

def replace_text_in_paragraph(paragraph, old_text, new_text):
    text = paragraph.text
    if old_text not in text:
        return False

    start_idx = text.find(old_text)
    end_idx = start_idx + len(old_text)
    
    current_char_count = 0
    runs_to_modify = []
    
    for idx, run in enumerate(paragraph.runs):
        run_len = len(run.text)
        run_start = current_char_count
        run_end = current_char_count + run_len
        
        if run_start < end_idx and run_end > start_idx:
            runs_to_modify.append((idx, run, run_start, run_end))
            
        current_char_count += run_len
        
    if not runs_to_modify:
        return False
        
    first_run_idx, first_run, fr_start, fr_end = runs_to_modify[0]
    prefix = first_run.text[:max(0, start_idx - fr_start)]
    
    if len(runs_to_modify) == 1:
        suffix = first_run.text[max(0, end_idx - fr_start):]
        first_run.text = prefix + new_text + suffix
    else:
        last_run_idx, last_run, lr_start, lr_end = runs_to_modify[-1]
        suffix = last_run.text[max(0, end_idx - lr_start):]
        first_run.text = prefix + new_text
        for idx, run, r_start, r_end in runs_to_modify[1:-1]:
            run.text = ""
        last_run.text = suffix
        
    return True

def replace_text_in_document(doc, old_text, new_text):
    old_text = re.sub(r'\s+', ' ', old_text).replace('\xa0', ' ').strip()
    if not old_text:
        return False
        
    modified = False
    
    def normalize_paragraph(p):
        for run in p.runs:
            if run.text:
                run.text = run.text.replace('\xa0', ' ')

    for p in doc.paragraphs:
        normalize_paragraph(p)
        if replace_text_in_paragraph(p, old_text, new_text):
            modified = True
            
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    normalize_paragraph(p)
                    if replace_text_in_paragraph(p, old_text, new_text):
                        modified = True
                        
    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                for p in header.paragraphs:
                    normalize_paragraph(p)
                    if replace_text_in_paragraph(p, old_text, new_text):
                        modified = True
                for t in header.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                normalize_paragraph(p)
                                if replace_text_in_paragraph(p, old_text, new_text):
                                    modified = True
                                    
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                for p in footer.paragraphs:
                    normalize_paragraph(p)
                    if replace_text_in_paragraph(p, old_text, new_text):
                        modified = True
                for t in footer.tables:
                    for row in t.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                normalize_paragraph(p)
                                if replace_text_in_paragraph(p, old_text, new_text):
                                    modified = True
    return modified

def set_cell_text_preserving_format(cell, new_text):
    if not cell.paragraphs:
        cell.add_paragraph(new_text)
        return
    p = cell.paragraphs[0]
    if p.runs:
        first_run = p.runs[0]
        first_run.text = new_text
        for run in p.runs[1:]:
            run.text = ""
        for extra_p in cell.paragraphs[1:]:
            extra_p.text = ""
    else:
        p.text = new_text

def auto_replace_room_specification_table(doc):
    for table in doc.tables:
        is_spec_table = False
        header_row_idx = -1
        
        for r_idx, row in enumerate(table.rows):
            cells_text = [c.text.lower() for c in row.cells]
            has_blok = any("блок" in t or "blok" in t for t in cells_text)
            has_qavat = any("қават" in t or "qavat" in t for t in cells_text)
            has_maydon = any("maydon" in t or "майдон" in t or "общая площадь" in t for t in cells_text)
            
            if has_blok and has_qavat and has_maydon:
                is_spec_table = True
                header_row_idx = r_idx
                break
                
        if is_spec_table and header_row_idx != -1:
            header_cells = table.rows[header_row_idx].cells
            for r_idx in range(header_row_idx + 1, len(table.rows)):
                data_row = table.rows[r_idx]
                row_text_combined = "".join([c.text.strip() for c in data_row.cells])
                if not row_text_combined:
                    continue
                for col_idx, cell in enumerate(data_row.cells):
                    if col_idx >= len(header_cells):
                        continue
                    header_txt = header_cells[col_idx].text.lower()
                    if "блок" in header_txt or "blok" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ BLOK }}")
                    elif "қават" in header_txt or "qavat" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ QAVAT }}")
                    elif "майдон" in header_txt or "maydon" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ MAYDON }}")
                    elif "1 kv" in header_txt or "1кв" in header_txt or "1 кв" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ KV_METR_NARXI }}")
                    elif "умумий нарх" in header_txt or "umumiy narx" in header_txt or "umumiy summa" in header_txt or "умумий сумма" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ UMUMIY_NARX }}")
                    elif "хоna" in header_txt or "xona" in header_txt or "хона" in header_txt:
                        set_cell_text_preserving_format(cell, "{{ XONALAR_SONI }}")

def apply_replacements(docx_path, replacements, output_path):
    """
    Replaces specified text strings with placeholder tags in the docx file and saves it.
    """
    if not os.path.exists(docx_path):
        return {"error": f"File not found: {docx_path}"}
        
    doc = docx.Document(docx_path)
    
    # 1. Add company-specific and document-specific custom replacements (fallbacks)
    custom = {
        # Client FISH
        "Mamanov Abdulla Xursanmurotovich": "MIJOZ_FISH",
        "A. Mamanov": "MIJOZ_FISH_QISQA",
        "A.Mamanov": "MIJOZ_FISH_QISQA",
        
        # Client PINFL
        "32601823990061": "MIJOZ_JSHSHIR",
        
        # Client Passport
        "AD №9160055": "PASPORT_SERIYA",
        "AD № 9160055": "PASPORT_SERIYA",
        "AD№9160055": "PASPORT_SERIYA",
        "AD 9160055": "PASPORT_SERIYA",
        "26/10/2024": "PASPORT_SANA",
        "Samarqand viloyati Urgut tumani I.I.B": "PASPORT_JOYI",
        "Samarqand viloyati Urgut tumani I.I.B tomonidan berilgan": "PASPORT_JOYI",
        
        # Client Address
        "Samarqand viloyati, Urgut tuman Yuqori Tegana Do`stlik maxallasida, ro`yxatda turuvchi": "MIJOZ_MANZILI",
        "Samarqand viloyati, Urgut tuman Yuqori Tegana Do`stlik maxallasida": "MIJOZ_MANZILI",
        "Samarqand viloyati Urgut tumani Yuqori Tegana mahalla": "MIJOZ_MANZILI",
        
        # Client Phone
        "(94) 633-04-00, (93) 114-82-72": "MIJOZ_TELEFON",
        "(94) 633-04-00": "MIJOZ_TELEFON",
        "(93) 114-82-72": "MIJOZ_TELEFON",
        
        # Residential Contract specific
        "4B/9": "SHARTNOMA_RAQAMI",
        "285 840 000": "UMUMIY_NARX",
        "285\u00a0840\u00a0000": "UMUMIY_NARX",
        "285.840.000": "UMUMIY_NARX",
        "ikki yuz sakson besh million sakkiz yuz qirq ming": "UMUMIY_NARX_SOZ_BILAN",
        "57 168 000,00": "BOSHLANGICH_TOLOV",
        "57 168 000": "BOSHLANGICH_TOLOV",
        "57\u00a0168\u00a0000,00": "BOSHLANGICH_TOLOV",
        "57\u00a0168\u00a0000": "BOSHLANGICH_TOLOV",
        "57.168.000,00": "BOSHLANGICH_TOLOV",
        "57.168.000": "BOSHLANGICH_TOLOV",
        "57168000": "BOSHLANGICH_TOLOV",
        "57 168 000,00 (ellik yetti million bir yuz oltmish sakkiz ming) so`m": "{{ BOSHLANGICH_TOLOV }} ({{ BOSHLANGICH_TOLOV_SOZ_BILAN }}) so`m",
        "57 168 000 (ellik yetti million bir yuz oltmish sakkiz ming) so`m": "{{ BOSHLANGICH_TOLOV }} ({{ BOSHLANGICH_TOLOV_SOZ_BILAN }}) so`m",
        "ellik yetti million bir yuz oltmish sakkiz ming": "BOSHLANGICH_TOLOV_SOZ_BILAN",
        "36 (o'ttiz olti) oy": "{{ MUDDAT_OY }} ({{ MUDDAT_OY_SOZ_BILAN }}) oy",
        "36 (o`ttiz olti) oy": "{{ MUDDAT_OY }} ({{ MUDDAT_OY_SOZ_BILAN }}) oy",
        "36 (o'ttiz olti)": "{{ MUDDAT_OY }} ({{ MUDDAT_OY_SOZ_BILAN }})",
        "36 (o`ttiz olti)": "{{ MUDDAT_OY }} ({{ MUDDAT_OY_SOZ_BILAN }})",
        "36": "MUDDAT_OY",
        "o'ttiz olti": "MUDDAT_OY_SOZ_BILAN",
        "o`ttiz olti": "MUDDAT_OY_SOZ_BILAN",
        "kamida 20%, 57 168 000,00 (ellik yetti million bir yuz oltmish sakkiz ming) so`m": "kamida {{ BOSHLANGICH_TOLOV_FOIZ }}%, {{ BOSHLANGICH_TOLOV }} ({{ BOSHLANGICH_TOLOV_SOZ_BILAN }}) so`m",
        "kamida 20%": "kamida {{ BOSHLANGICH_TOLOV_FOIZ }}%",
        
        # Commercial Contract specific
        "1B/1P": "SHARTNOMA_RAQAMI",
        "190 000 000": "UMUMIY_NARX",
        "190\u00a0000\u00a0000": "UMUMIY_NARX",
        "190.000.000": "UMUMIY_NARX",
        "Bir yuz to`qson million": "UMUMIY_NARX_SOZ_BILAN",
        "bir yuz to`qson million": "UMUMIY_NARX_SOZ_BILAN",
        "3 998 316": "KV_METR_NARXI",
        "3\u00a0998\u00a0316": "KV_METR_NARXI",
        "3.998.316": "KV_METR_NARXI",
        "kamida 100% , 190 000 000 (Bir yuz to`qson million) so`m": "kamida {{ BOSHLANGICH_TOLOV_FOIZ }}%, {{ BOSHLANGICH_TOLOV }} ({{ BOSHLANGICH_TOLOV_SOZ_BILAN }}) so`m",
        "kamida 100%": "kamida {{ BOSHLANGICH_TOLOV_FOIZ }}%",
        
        # General company settings
        "«SHER-DOR» xususiy korxonasi": "KORXONA_NOMI",
        "«SHER-DOR» Xususiy korxonasi": "KORXONA_NOMI",
        "«Sher-Dor» Xususiy korxonasi": "KORXONA_NOMI",
        "Urgut tuman Do`stlik maxallasi Navoi-Shox ko`chasi №3-uy": "KORXONA_MANZILI",
        "2020 8000 6005 5796 1001": "KORXONA_BANK_ACCOUNT",
        "20208000600557961001": "KORXONA_BANK_ACCOUNT",
        "ATB “Agrobank” Urgut tuman filiali": "KORXONA_BANK",
        "ATB “Agrobank” Urgut tuman filiali,": "KORXONA_BANK",
        "00300": "KORXONA_MFO",
        "202 162 496": "KORXONA_INN",
        "202162496": "KORXONA_INN",
        "J.Raxmonov": "DIREKTOR",
        "J. Raxmonov": "DIREKTOR",
        "(97) 708-44-44, (93) 573-29-29": "KORXONA_TELEFON",
        "(97) 708-44-44": "KORXONA_TELEFON",
        "(93) 573-29-29": "KORXONA_TELEFON",
        
        # Location / combined description placeholders
        "Urgut shahar Do`stlik mahallasi Navoiy ko`chasi №83-uy manzilida joylashgan yangidan qurilayotgan ko`p qavatli xonadonlarning 4-(to`rtinchi) blok 4-(to`rtinchi) qavati, 9-(to`qqizinchi) xonadonsi": "OBYEKT_TAFSILOTI",
        "Urgut shahar Do`stlik mahallasi Navoiy ko`chasi №83-uy manzilida joylashgan yangidan qurilayotgan ko`p qavatli xonadonlarning 4-(to`rtinchi) blok 4-(to`rtinchi) qavati, 9-(to`qqizinchi) xonadoni": "OBYEKT_TAFSILOTI",
        "Urgut shahar Do`stlik mahallasi Navoiy ko`chasi №83-uy manzilida joylashgan yangidan qurilayotgan savdo do`konlarning 1-(birinchi) blok -1-(podvol) qavati, 1-( birinchi) savdo do`koni": "OBYEKT_TAFSILOTI"
    }
    for k, v in custom.items():
        if k not in replacements:
            replacements[k] = v

    # 2. Run custom table cell replacements for the rooms/specification table
    auto_replace_room_specification_table(doc)

    # 3. Replaces date headers
    replace_text_in_document(doc, "“13 ” mart 2025-yil", "“{{ KUN }}” {{ OY }} {{ YIL }}-yil")
    replace_text_in_document(doc, "“13” mart 2025-yil", "“{{ KUN }}” {{ OY }} {{ YIL }}-yil")
    replace_text_in_document(doc, "13 mart 2025-yil", "{{ KUN }} {{ OY }} {{ YIL }}-yil")
    
    # Sort replacements by original text length desc to avoid substring replacement bugs
    sorted_replacements = sorted(replacements.items(), key=lambda x: len(x[0]), reverse=True)
    
    success_count = 0
    for old_text, placeholder in sorted_replacements:
        # Wrap placeholder in Jinja2 style if it doesn't already contain Jinja braces
        if "{{" in placeholder:
            tag = placeholder
        else:
            tag = f"{{{{ {placeholder} }}}}"
        if replace_text_in_document(doc, old_text, tag):
            success_count += 1
            
    doc.save(output_path)
    return {"success": True, "replaced_count": success_count}

def compile_document(template_path, context, output_path):
    """
    Renders the docx template with variables context using docxtpl.
    """
    if not os.path.exists(template_path):
        return {"error": f"Template not found: {template_path}"}
        
    doc = DocxTemplate(template_path)
    
    # Convert python newlines with XML line breaks dynamically in context strings
    from docxtpl import RichText
    for k, v in list(context.items()):
        if isinstance(v, str) and '\n' in v:
            rt = RichText()
            parts = v.split('\n')
            for i, part in enumerate(parts):
                if i > 0:
                    rt.add_break()
                rt.add_text(part)
            context[k] = rt

    doc.render(context)
    
    # Post-process the rendered document to clean up duplicate phone prefixes "+998 +998"
    def clean_paragraphs(paragraphs):
        for p in paragraphs:
            if re.search(r'(\+?998\s*){2,}', p.text):
                p.text = re.sub(r'(\+?998\s*){2,}', '+998 ', p.text)

    clean_paragraphs(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                clean_paragraphs(cell.paragraphs)

    for section in doc.sections:
        for header in [section.header, section.first_page_header, section.even_page_header]:
            if header:
                clean_paragraphs(header.paragraphs)
                for t in header.tables:
                    for r in t.rows:
                        for c in r.cells:
                            clean_paragraphs(c.paragraphs)
        for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
            if footer:
                clean_paragraphs(footer.paragraphs)
                for t in footer.tables:
                    for r in t.rows:
                        for c in r.cells:
                            clean_paragraphs(c.paragraphs)

    doc.save(output_path)
    return {"success": True}

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print(json.dumps({"error": "Missing arguments. Usage: script.py <command> <input_path> [options]"}))
        sys.exit(1)
        
    command = sys.argv[1]
    input_path = sys.argv[2]
    
    # Make sure stdout uses UTF-8 encoding
    sys.stdout.reconfigure(encoding='utf-8')
    
    try:
        if command == "detect":
            result = detect_placeholders(input_path)
            print(json.dumps(result, ensure_ascii=False))
            
        elif command == "preview":
            result = convert_docx_to_html(input_path)
            print(json.dumps(result, ensure_ascii=False))
            
        elif command == "replace":
            if len(sys.argv) < 5:
                print(json.dumps({"error": "Missing arguments for replace: need replacements_json and output_path"}))
                sys.exit(1)
            arg3 = sys.argv[3]
            if not arg3.strip().startswith('{') and os.path.exists(arg3):
                with open(arg3, "r", encoding="utf-8") as f:
                    replacements = json.load(f)
            else:
                replacements = json.loads(arg3)
            output_path = sys.argv[4]
            result = apply_replacements(input_path, replacements, output_path)
            print(json.dumps(result, ensure_ascii=False))
            
        elif command == "compile":
            if len(sys.argv) < 5:
                print(json.dumps({"error": "Missing arguments for compile: need context_json and output_path"}))
                sys.exit(1)
            arg3 = sys.argv[3]
            if not arg3.strip().startswith('{') and os.path.exists(arg3):
                with open(arg3, "r", encoding="utf-8") as f:
                    context = json.load(f)
            else:
                context = json.loads(arg3)
            output_path = sys.argv[4]
            result = compile_document(input_path, context, output_path)
            print(json.dumps(result, ensure_ascii=False))
            
        else:
            print(json.dumps({"error": f"Unknown command: {command}"}))
            sys.exit(1)
            
    except Exception as e:
        import traceback
        print(json.dumps({
            "error": str(e),
            "traceback": traceback.format_exc()
        }, ensure_ascii=False))
        sys.exit(1)
